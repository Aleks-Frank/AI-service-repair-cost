"""Парсер прайс-листов поставщиков материалов: выгрузки из 1С (xls/xlsx) и каталоги магазинов (json).

Как и парсер ПЗР, ничего не знает о базе: на входе файл, на выходе PriceListResult.

Формат JSON (каталоги «Строительный двор», «Кувалда», «Лемана Про»):
    {"generation_date", "shop": {...}, "products": [{"article_code", "vendor_code"?, "name",
     "breadcrumbs"?, "brand"?, "specifications"?, "cities": [{"city_name",
     "price": {"date", "price"}, "stock"?: [{"type": "all", "count"}, ...]}]}]}
    Единиц измерения в каталогах нет — цена считается за штуку; остаток берётся из записи
    «all», если магазин его передаёт, иначе остаток неизвестен.

Ожидаемая структура xls (как в «прайс для сервисной компании.xls» от «Р.О.С. электро»):
    шапка    — название поставщика («ООО "…"») и строка «Цены указаны на ДД.ММ.ГГГГ»;
    заголовок — строка с колонкой «…Артикул»; колонки ищутся по словам, а не по буквам:
               «Номенклатура», «Артикул», «Остаток», «Распродажа», «Цена», «Ед.»;
    группы   — строки без цены; вложенность задаёт уровень группировки строк Excel
               (outline level): 0 — бренд («ИЕК», «ITK»), 1 и глубже — категории;
    позиции  — строки с ценой: наименование, артикул, остаток, цена, единица.
Другой поставщик с другими заголовками колонок потребует дописать слова в _COLUMNS.
"""

import json
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from rag.normalizers import clean_text, normalize_unit, to_decimal

# Колонки прайса: ключ -> (слова, одно из которых должно быть в заголовке; слова-исключения).
# Порядок важен: «Номенклатура.Артикул» содержит и «номенклатура», поэтому артикул ищем первым
# и исключаем его колонку при поиске наименования.
_COLUMNS: dict[str, tuple[tuple[str, ...], tuple[str, ...]]] = {
    "article": (("артикул",), ()),
    "name": (("номенклатура", "наименование"), ("артикул",)),
    "stock": (("остаток",), ()),
    "sale": (("распродажа",), ()),
    "price": (("цена",), ()),
    "unit": (("ед.", "ед ", "единица"), ()),
}
_REQUIRED_COLUMNS = ("article", "name", "price", "unit")

# Организационно-правовая форма в начале строки шапки — признак названия поставщика.
_SUPPLIER_RE = re.compile(r"^(ООО|АО|ПАО|ЗАО|ОАО|ИП)\b")
_DATE_RE = re.compile(r"(\d{2})\.(\d{2})\.(\d{4})")
# Фасовка в наименовании: «(8м)», «(120м)», «(4шт/компл)», «(20шт)».
_PACK_RE = re.compile(r"\((\d+(?:[.,]\d+)?)\s*(м|шт)(?=[\s/)])")

# Частые сокращения в наименованиях прайса. Модели эмбеддингов и reranker их не понимают
# («Роз. с з/к о/у» не похоже на «розетка с заземлением открытой установки»), поэтому в тексте
# для поиска они раскрываются. Наименование в базе остаётся как в прайсе.
_ABBREVIATIONS = {
    "Роз.": "Розетка", "Розет.": "Розетка", "Выключ.": "Выключатель", "Выкл.": "Выключатель",
    "з/к": "заземляющий контакт", "з/ш": "защитные шторки",
    "о/у": "открытой установки", "с/у": "скрытой установки", "откр. уст.": "открытой установки",
    "1-кл.": "одноклавишный", "2-кл.": "двухклавишный", "3-кл.": "трехклавишный",
    "1-мест.": "одноместная", "2-мест.": "двухместная", "3-мест.": "трехместная", "4-мест.": "четырехместная",
    "прох.": "проходной", "перекр.": "перекрестный", "перенос.": "переносная", "встр.": "встраиваемая",
    "стац.": "стационарная", "расп.": "распределительная", "гофр.": "гофрированная", "сб.": "сборе",
    "бел.": "белый", "черн.": "черный", "граф.": "графит", "крем.": "кремовый", "Пож.": "Пожарный",
}
# Сокращение заменяется только целиком: слева не буква/цифра/дефис, справа не буква.
_ABBREVIATION_RE = re.compile(
    r"(?<![\w-])(" + "|".join(re.escape(key) for key in sorted(_ABBREVIATIONS, key=len, reverse=True)) + r")(?!\w)"
)


def expand_abbreviations(text: str) -> str:
    """Раскрывает частые сокращения прайса: «Роз. с з/к о/у» -> «Розетка с заземляющий контакт открытой установки».

    Принимает: text — наименование позиции.
    Возвращает: текст с раскрытыми сокращениями (для эмбеддинга и полнотекстового поиска).
    """
    return _ABBREVIATION_RE.sub(lambda match: _ABBREVIATIONS[match.group(1)], text)


# --- Сопоставление одного товара у разных поставщиков ---------------------------------

# Кириллические буквы, которые при наборе артикула легко спутать с латинскими.
# «СКМР10D» (часть букв кириллицей) и «CKMP10D» — один артикул.
_LOOKALIKE = str.maketrans("АВЕКМНОРСТУХ", "ABEKMHOPCTYX")


def normalize_article(article: str) -> str:
    """Приводит артикул к виду для сравнения между поставщиками.

    Верхний регистр, без пробелов, похожие кириллические буквы -> латинские.
    Суффиксы («-R», «-K01») не отрезаются: это могут быть разные модификации товара.

    Принимает: article — артикул из прайса.
    Возвращает: нормализованный артикул.
    """
    return re.sub(r"\s+", "", article).upper().translate(_LOOKALIKE)


def pricing_units(unit: str, pack_size: Decimal | None, pack_unit: str | None) -> tuple[str, Decimal, bool]:
    """Определяет базовую единицу, сколько её в единице цены и обязательна ли покупка упаковками.

    Правило: кратность обязательна, только если цена указана за упаковку («компл», «уп»).
        «компл» + «(4шт/компл)» -> базовая «шт», 4 шт в единице цены, покупка комплектами;
        «уп» без фасовки в названии -> базовая «уп», состав упаковки неизвестен;
        «шт» + «(6шт)», «м» + «(8м)» -> цена уже за штуку/метр, коробка справочно.

    Принимает: unit — единица цены; pack_size, pack_unit — фасовка из наименования.
    Возвращает: (базовая единица, базовых единиц в единице цены, покупать ли целыми упаковками).
    """
    if unit in {"компл", "уп"}:
        if pack_size and pack_unit:
            return pack_unit, pack_size, True
        return unit, Decimal(1), True
    return unit, Decimal(1), False


# Ключевые параметры товара из наименования. Если у эталона есть такой параметр, у товара
# другого поставщика он должен совпасть — иначе «Угол 40х16» спутается с «Углом 40х25».
_PARAM_PATTERNS: dict[str, re.Pattern[str]] = {
    "size": re.compile(r"(?<!\d)(\d+(?:[.,]\d+)?(?:\s*[хxХX×*/]\s*\d+(?:[.,]\d+)?)+)(?!\d)"),
    "current": re.compile(r"(?<![\w.,])(\d+(?:[.,]\d+)?)\s?[АA](?![\w])"),
    "power": re.compile(r"(?<![\w.,])(\d+(?:[.,]\d+)?)\s?Вт(?![\w])"),
    "color_temp": re.compile(r"(?<![\w.,])(\d{4})\s?[КK](?![\w])"),
    "socket": re.compile(r"(?<![\w])([ЕE]\d{2}|G[UX]?\d{1,2}(?:[.,]\d)?)(?![\w])"),
    "ip": re.compile(r"(?<![\w])IP\s?(\d{2})(?![\w])"),
}
_COLORS = {
    "бел": "белый", "черн": "черный", "сер": "серый", "крем": "кремовый", "беж": "бежевый",
    "граф": "графит", "син": "синий", "красн": "красный", "желт": "желтый", "зелен": "зеленый",
    "коричн": "коричневый", "алюм": "алюминий", "сереб": "серебро", "золот": "золото",
}
# Исполнение, которое отличает разные товары с одинаковыми размерами: «Угол внутренний 40х40»
# и «Угол внешний 40х40». «Наружный» — синоним «внешнего» (так пишут у разных поставщиков).
_KINDS = {"внутрен": "внутренний", "внешн": "внешний", "наружн": "внешний"}
# Параметры, без которых товар считается другим (если они есть у эталона, а у кандидата нет).
CRITICAL_PARAMS = ("size", "current", "power", "socket")


def extract_params(name: str) -> dict[str, list[str]]:
    """Извлекает ключевые параметры из наименования: размер, ток, мощность, цоколь, IP, цвет.

    Принимает: name — наименование позиции.
    Возвращает: словарь {параметр: отсортированный список значений}; пример:
        «Угол внешний КМН 40х16 белый (4шт/компл)» -> {"size": ["40x16"], "color": ["белый"], "kind": ["внешний"]}.
    """
    params: dict[str, list[str]] = {}
    for key, pattern in _PARAM_PATTERNS.items():
        values = set()
        for match in pattern.finditer(name):
            # Единый вид значения: латиница, точка вместо запятой, «х» -> «x», без пробелов.
            value = re.sub(r"\s+", "", match.group(1)).upper().translate(_LOOKALIKE).replace(",", ".")
            values.add(re.sub(r"[XХ×*]", "x", value) if key == "size" else value)
        if values:
            params[key] = sorted(values)
    words = re.findall(r"[а-яё]+", name.lower())
    colors = {color for stem, color in _COLORS.items() for word in words if word.startswith(stem)}
    if colors:
        params["color"] = sorted(colors)
    kinds = {kind for stem, kind in _KINDS.items() for word in words if word.startswith(stem)}
    if kinds:
        params["kind"] = sorted(kinds)
    return params


# Обозначение модели: слово с латинскими буквами («Colour», «Como», «C35», «VTM.303»).
_MODEL_TOKEN_RE = re.compile(r"(?<![\w])[A-Za-z0-9][A-Za-z0-9.\-/]*[A-Za-z][A-Za-z0-9.\-/]*|(?<![\w])[A-Za-z][A-Za-z0-9.\-/]*")
# Латинские слова, которые не отличают модели: единицы, цоколи, степень защиты, общие слова.
_NOT_MODEL = {"led", "ip", "w", "v", "k", "mm", "cm", "rgb", "white", "black", "pro", "max", "mini"}


def model_tokens(name: str, brand: str | None) -> set[str]:
    """Обозначения модели в наименовании — латинские слова, кроме бренда, цоколя и общих слов.

    Пример: «Раковина Cersanit Colour керамика 50 см» (бренд Cersanit) -> {"colour"}.

    Принимает: name — наименование; brand — бренд позиции (может быть None).
    Возвращает: множество обозначений в нижнем регистре.
    """
    params = extract_params(name)
    skip = {value.lower() for key in ("socket", "ip") for value in params.get(key, [])} | _NOT_MODEL
    if brand:
        skip |= {word.lower() for word in re.findall(r"[A-Za-z0-9]+", brand)}
    tokens = {token.lower().strip(".-/") for token in _MODEL_TOKEN_RE.findall(name)}
    return {token for token in tokens if token and token not in skip and not token.startswith("ip")}


def product_type(name: str) -> str | None:
    """Тип изделия — первое русское слово наименования (без аббревиатур в верхнем регистре).

    «РКС-20-30-П-К Розетка с з/к» -> «розет»; «Разъем РШ-ВШ 32-002» -> «разъе»; «VEGA Роз. …» -> «розет».
    Сравниваются первые 5 букв, чтобы «розетка» и «розетки» совпадали.

    Принимает: name — наименование позиции.
    Возвращает: начало слова-типа в нижнем регистре или None.
    """
    for word in re.findall(r"[А-ЯЁа-яё]{3,}", expand_abbreviations(name)):
        if not word.isupper():
            return word.lower().replace("ё", "е")[:5]
    return None


def models_match(reference: set[str], candidate: set[str]) -> bool:
    """Проверяет, что обозначения моделей не противоречат друг другу.

    Если у обоих товаров есть обозначения модели, нужно хотя бы одно общее: «Colour» и «Como»
    одного бренда — разные раковины. Если у одного из товаров обозначений нет, проверка не
    применяется (название могло быть сокращено).

    Принимает: reference, candidate — результаты model_tokens.
    Возвращает: True, если товары можно считать одной моделью.
    """
    return not (reference and candidate) or bool(reference & candidate)


def params_match(reference: dict[str, list[str]], candidate: dict[str, list[str]]) -> bool:
    """Проверяет, что у кандидата те же ключевые параметры, что и у эталона.

    Правила: параметр, который есть у обоих, должен совпасть; критичный параметр эталона
    (размер, ток, мощность, цоколь) обязан быть и у кандидата. Цвет или IP, не указанные
    у кандидата, не считаются расхождением — их часто опускают в названии.

    Принимает: reference, candidate — результаты extract_params.
    Возвращает: True, если товары можно считать одинаковыми по параметрам.
    """
    for key, values in reference.items():
        if key in candidate:
            if candidate[key] != values:
                return False
        elif key in CRITICAL_PARAMS:
            return False
    return True


@dataclass
class ParsedMaterial:
    """Одна позиция прайс-листа."""

    article: str
    name: str
    brand: str | None
    category: str
    unit: str
    unit_raw: str
    price: Decimal
    stock: Decimal | None
    on_sale: bool
    pack_size: Decimal | None
    pack_unit: str | None
    source_row: int
    # Характеристики из каталога магазина («Цвет: белый; Диаметр: 25 мм»); у прайсов 1С пусто.
    details: str = ""

    @property
    def search_text(self) -> str:
        """Текст для поиска: наименование с раскрытыми сокращениями, путь категории и характеристики.

        Цена и поставщик не входят: они не описывают сам материал.

        Принимает: ничего.
        Возвращает: строку для модели эмбеддингов и reranker'а.
        """
        parts = [expand_abbreviations(self.name)]
        if self.category:
            parts.append(f"Категория: {self.category}")
        if self.details:
            parts.append(f"Характеристики: {self.details}")
        return "\n".join(parts)

    @property
    def pricing(self) -> tuple[str, Decimal, bool]:
        """Базовая единица, базовых единиц в единице цены и обязательность покупки упаковками.

        Принимает: ничего.
        Возвращает: кортеж из pricing_units().
        """
        return pricing_units(self.unit, self.pack_size, self.pack_unit)

    @property
    def base_price(self) -> Decimal:
        """Цена за одну базовую единицу (за 1 шт или 1 м) — по ней сравниваются поставщики.

        Принимает: ничего.
        Возвращает: цену с точностью до 0,0001 руб.
        """
        return (self.price / self.pricing[1]).quantize(Decimal("0.0001"))


@dataclass
class PriceListResult:
    """Результат разбора прайс-листа."""

    source_file: str
    supplier_name: str | None
    price_date: date | None
    materials: list[ParsedMaterial] = field(default_factory=list)
    # Пропущенные строки: (номер строки, причина) — без цены, без артикула, дубль артикула.
    skipped: list[tuple[int, str]] = field(default_factory=list)

    def summary(self) -> dict:
        """Сводка для печати: сколько позиций, категорий, единиц и пропусков.

        Принимает: ничего.
        Возвращает: словарь со счётчиками.
        """
        return {
            "source_file": self.source_file,
            "supplier_name": self.supplier_name,
            "price_date": self.price_date.isoformat() if self.price_date else None,
            "materials": len(self.materials),
            "categories": len({m.category for m in self.materials}),
            "units": dict(Counter(m.unit for m in self.materials).most_common()),
            "with_pack": sum(m.pack_size is not None for m in self.materials),
            # Позиции, которые покупаются только целыми упаковками (цена за «компл»/«уп»).
            "sold_by_pack": sum(m.pricing[2] for m in self.materials),
            "skipped": dict(Counter(reason for _, reason in self.skipped)),
        }


def _read_rows(path: Path) -> list[tuple[int, int, list]]:
    """Читает первый лист файла вместе с уровнями группировки строк.

    Принимает: path — путь к .xls или .xlsx.
    Возвращает: список (номер строки с 1, уровень группировки, значения ячеек).
    """
    if path.suffix.lower() == ".xls":
        import xlrd

        # formatting_info=True нужен, чтобы xlrd прочитал уровни группировки строк (rowinfo_map).
        sheet = xlrd.open_workbook(path, formatting_info=True).sheet_by_index(0)
        return [
            (index + 1, sheet.rowinfo_map[index].outline_level if index in sheet.rowinfo_map else 0, sheet.row_values(index))
            for index in range(sheet.nrows)
        ]

    from openpyxl import load_workbook

    # Обычный режим (не read_only): только в нём доступны row_dimensions с уровнями группировки.
    sheet = load_workbook(path, data_only=True).worksheets[0]
    return [
        (index, sheet.row_dimensions[index].outline_level or 0, list(values))
        for index, values in enumerate(sheet.iter_rows(values_only=True), start=1)
    ]


def _find_header(rows: list[tuple[int, int, list]]) -> tuple[int, dict[str, int]]:
    """Находит строку заголовка и номера колонок по словам в заголовках.

    Заголовок бывает в две строки («6. для тендеров» над «Цена»), поэтому текст колонки
    склеивается из строки с «Артикул» и следующей за ней.

    Заголовок — первая строка со словом «артикул», в которой нашлись все обязательные колонки
    (слово может встретиться и в шапке, например в названии поставщика).

    Принимает: rows — строки листа.
    Возвращает: (индекс строки заголовка в rows, {ключ колонки: номер колонки}).
    Исключения: ValueError, если заголовок или обязательные колонки не найдены.
    """
    missing: list[str] = []
    for index, (_, _, values) in enumerate(rows):
        texts = [(clean_text(value) or "").lower() for value in values]
        if not any("артикул" in text for text in texts):
            continue
        below = rows[index + 1][2] if index + 1 < len(rows) else []
        headers = [f"{text} {(clean_text(below[col]) or '').lower() if col < len(below) else ''} "
                   for col, text in enumerate(texts)]
        columns: dict[str, int] = {}
        for key, (words, excluded) in _COLUMNS.items():
            # next() берёт первую подходящую колонку, которую ещё не заняла другая.
            column = next(
                (col for col, header in enumerate(headers)
                 if col not in columns.values()
                 and any(word in header for word in words)
                 and not any(word in header for word in excluded)),
                None,
            )
            if column is not None:
                columns[key] = column
        missing = [key for key in _REQUIRED_COLUMNS if key not in columns]
        if not missing:
            return index, columns
    if missing:
        raise ValueError(f"В заголовке прайса не найдены колонки: {', '.join(missing)}")
    raise ValueError("Не найдена строка заголовка с колонкой «Артикул».")


def _header_info(rows: list[tuple[int, int, list]]) -> tuple[str | None, date | None]:
    """Достаёт из шапки название поставщика и дату цен.

    Принимает: rows — строки шапки (до заголовка таблицы).
    Возвращает: (название или None, дата или None).
    """
    supplier, price_date = None, None
    for _, _, values in rows:
        for value in values:
            text = clean_text(value)
            if not text:
                continue
            if supplier is None and _SUPPLIER_RE.match(text):
                supplier = text
            match = _DATE_RE.search(text)
            if price_date is None and match and "цен" in text.lower():
                day, month, year = map(int, match.groups())
                price_date = date(year, month, day)
    return supplier, price_date


def _cell_text(value: object) -> str | None:
    """Текст ячейки; числа без «.0» (артикул 12345 xlrd отдаёт как 12345.0).

    Принимает: value — значение ячейки.
    Возвращает: строку или None.
    """
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return clean_text(value)


def _pack(name: str) -> tuple[Decimal | None, str | None]:
    """Извлекает фасовку из наименования: «Кабель-канал 16х16 (84м)» -> (84, «м»).

    Принимает: name — наименование позиции.
    Возвращает: (количество, единица) или (None, None).
    """
    match = _PACK_RE.search(name)
    if not match:
        return None, None
    return to_decimal(match.group(1)), match.group(2)


def parse_price_list(path: str | Path) -> PriceListResult:
    """Разбирает прайс-лист поставщика: выгрузку из 1С (.xls/.xlsx) или каталог магазина (.json).

    Принимает: path — путь к файлу.
    Возвращает: PriceListResult с позициями и пропущенными строками.
    """
    path = Path(path)
    if path.suffix.lower() == ".json":
        return _parse_shop_json(path)
    return _parse_spreadsheet(path)


def _parse_spreadsheet(path: Path) -> PriceListResult:
    """Разбирает выгрузку прайса из 1С (структура описана в начале модуля).

    Принимает: path — путь к .xls/.xlsx.
    Возвращает: PriceListResult.
    """
    rows = _read_rows(path)
    header_index, columns = _find_header(rows)
    supplier, price_date = _header_info(rows[:header_index])
    result = PriceListResult(source_file=str(path), supplier_name=supplier, price_date=price_date)

    def cell(values: list, key: str) -> object:
        """Значение колонки key в строке (None, если колонки нет в прайсе)."""
        column = columns.get(key)
        return values[column] if column is not None and column < len(values) else None

    # Текущие группы по уровням: {уровень: название}. Уровень 0 — бренд.
    groups: dict[int, str] = {}
    seen_articles: set[str] = set()
    # Данные начинаются после двухстрочного заголовка.
    for row_number, level, values in rows[header_index + 2:]:
        name = clean_text(cell(values, "name"))
        if not name:
            continue
        price = to_decimal(cell(values, "price"))
        unit_raw = clean_text(cell(values, "unit"))

        # --- Строка группы: запоминаем на её уровне и забываем более глубокие уровни. ---
        if price is None and not unit_raw:
            groups = {lvl: title for lvl, title in groups.items() if lvl < level}
            groups[level] = name
            continue

        # --- Позиция прайса. ---
        article = _cell_text(cell(values, "article"))
        if price is None:
            result.skipped.append((row_number, "no_price"))
            continue
        if not article:
            result.skipped.append((row_number, "no_article"))
            continue
        if article in seen_articles:
            result.skipped.append((row_number, "duplicate_article"))
            continue
        seen_articles.add(article)

        # Путь категории — группы выше позиции, кроме бренда; соседние повторы (в 1С группа
        # и подгруппа бывают названы одинаково) склеиваем через dict.fromkeys.
        path_titles = [title for lvl, title in sorted(groups.items()) if 0 < lvl < level]
        pack_size, pack_unit = _pack(name)
        result.materials.append(ParsedMaterial(
            article=article,
            name=name,
            brand=groups.get(0),
            category=" / ".join(dict.fromkeys(path_titles)),
            unit=normalize_unit(unit_raw),
            unit_raw=unit_raw or "",
            price=price,
            # Колонка «Остаток» есть, но ячейка пуста — товара нет на складе (0).
            # Колонки нет вовсе — остаток неизвестен (None).
            stock=(to_decimal(cell(values, "stock")) or Decimal(0)) if "stock" in columns else None,
            # В колонке «Распродажа» стоит «***», если позиция продаётся только из наличия.
            on_sale=bool(clean_text(cell(values, "sale"))),
            pack_size=pack_size,
            pack_unit=pack_unit,
            source_row=row_number,
        ))
    return result


# Служебные пункты «хлебных крошек» каталога, которые не являются категориями.
_BREADCRUMB_SKIP = {"главная", "каталог"}
# Сколько символов характеристик включать в текст для поиска (длинные описания только мешают).
_DETAILS_LIMIT = 400


def _city_offer(cities: list[dict]) -> tuple[Decimal | None, Decimal | None]:
    """Выбирает цену и остаток из записей по городам: запись с ценой и самой свежей датой цены.

    Принимает: cities — список записей товара по городам из каталога.
    Возвращает: (цена или None, остаток или None — если магазин остатки не передаёт).
    """
    priced = [c for c in cities if (c.get("price") or {}).get("price")]
    if not priced:
        return None, None
    best = max(priced, key=lambda c: c["price"].get("date") or "")
    stock_rows = best.get("stock") or []
    total = next((s.get("count") for s in stock_rows if s.get("type") == "all"), None)
    if total is None and stock_rows:
        total = sum(s.get("count") or 0 for s in stock_rows)
    return to_decimal(best["price"]["price"]), to_decimal(total) if stock_rows else None


def _parse_shop_json(path: Path) -> PriceListResult:
    """Разбирает каталог интернет-магазина в формате JSON (структура описана в начале модуля).

    Артикул — vendor_code, если он есть, иначе код товара в магазине (article_code). Это коды
    магазина, а не производителя, поэтому такие каталоги загружаются с флагом --own-articles.
    Категория — «хлебные крошки» без служебных пунктов и бренда.

    Принимает: path — путь к .json.
    Возвращает: PriceListResult.
    """
    with path.open(encoding="utf-8") as file:
        data = json.load(file)
    generated = data.get("generation_date")
    price_date = datetime.fromisoformat(generated.replace("Z", "+00:00")).date() if generated else None
    result = PriceListResult(
        source_file=str(path),
        supplier_name=(data.get("shop") or {}).get("shop_name"),
        price_date=price_date,
    )
    seen_articles: set[str] = set()
    for number, product in enumerate(data.get("products") or [], start=1):
        name = clean_text(product.get("name"))
        article = _cell_text(product.get("vendor_code")) or _cell_text(product.get("article_code"))
        if not name:
            continue
        # Магазин дублирует свой код в конце названия: «Герметик Tytan, белый [тов-213048]».
        if article:
            name = clean_text(name.replace(f"[{article}]", "")) or name
        price, stock = _city_offer(product.get("cities") or [])
        if price is None:
            result.skipped.append((number, "no_price"))
            continue
        if not article:
            result.skipped.append((number, "no_article"))
            continue
        if article in seen_articles:
            result.skipped.append((number, "duplicate_article"))
            continue
        seen_articles.add(article)

        brand = clean_text(product.get("brand"))
        crumbs = [clean_text(c) for c in product.get("breadcrumbs") or []]
        category = " / ".join(dict.fromkeys(
            c for c in crumbs if c and c.lower() not in _BREADCRUMB_SKIP and c != brand
        ))
        specs = product.get("specifications") or {}
        details = "; ".join(f"{clean_text(k)}: {clean_text(v)}" for k, v in specs.items() if clean_text(v))
        pack_size, pack_unit = _pack(name)
        result.materials.append(ParsedMaterial(
            article=article,
            name=name,
            brand=brand,
            category=category,
            # Единиц измерения в каталогах нет: цена магазина — за штуку товара.
            unit="шт",
            unit_raw="",
            price=price,
            stock=stock,
            on_sale=False,
            pack_size=pack_size,
            pack_unit=pack_unit,
            source_row=number,
            details=details[:_DETAILS_LIMIT],
        ))
    return result
